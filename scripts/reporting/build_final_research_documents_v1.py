from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "artifacts" / "final_research_report_v1"
FIGDIR = OUT / "figures"
FIGS = {int(re.search(r"Fig(\d+)", p.name).group(1)): p for p in FIGDIR.glob("Fig*.png")}
REPORT = OUT / "01_完整研究工作汇报.docx"
OUTLINE = OUT / "02_汇报PPT详细提纲.docx"
GLOSSARY_DOC = OUT / "03_术语与符号说明.docx"

NAVY, BLUE, TEAL, GOLD, RED, GRAY = "122D4F", "1A6EA8", "0E817A", "E2A32E", "BE3237", "6E7A86"
LIGHT, PALE_GOLD, WHITE = "EFF4F7", "FBF4E2", "FFFFFF"


TERMS = [
    ("MPC", "模型预测控制", "Model Predictive Control", "根据当前状态、预测信号与约束滚动求解有限时域优化，只执行首步动作。"),
    ("H1", "单步 MPC", "One-step MPC", "只考虑当前控制步的优化器或基线。"),
    ("H4", "四步 MPC", "Four-step MPC", "使用四个未来 15 分钟预测步；修复后覆盖至 t+60。"),
    ("Oracle", "特权未来真值", "Oracle future information", "只能用于离线教师或上界分析，不能进入学生在线输入。"),
    ("BC", "行为克隆", "Behavior Cloning", "把专家状态—动作对作为监督样本训练神经策略。"),
    ("RL", "强化学习", "Reinforcement Learning", "通过环境交互优化累计回报；本阶段未重新训练。"),
    ("SAC", "软演员评论家", "Soft Actor-Critic", "项目历史上使用离散变体评估 BC 初始化能否改善在线学习。"),
    ("Transformer", "Transformer 预测器", "Transformer forecaster", "编码 96 步历史并预测未来 4 步、4 类负载。"),
    ("Forecast Dataset v1", "预测数据集第一版", "Forecast Dataset v1", "只取已验证唯一 49 天周期，按时间切分且 scaler 只拟合训练集。"),
    ("Expert Dataset v2", "专家数据集第二版", "Expert Dataset v2", "由 repaired H4 Oracle 产生标签，学生只保留 34 维可部署观测。"),
    ("Triggered MPC", "事件触发 MPC", "Triggered MPC", "依据可部署风险分数在 H1 与 H4 之间选择。"),
    ("Adapter", "适配器", "Adapter", "将 SustainCluster 字段映射为统一状态、任务与动作语义的接口层。"),
    ("Information Contract", "信息契约", "Information Contract", "明确环境真值、控制器可见字段、预测量和禁止泄漏字段。"),
    ("Deployable", "可部署信息", "Deployable information", "在线决策时真实可获得或可由外部服务提供的信息。"),
    ("Privileged Knowledge", "特权知识", "Privileged knowledge", "训练期教师可使用、部署期学生不可直接观测的信息。"),
    ("Control Authority", "控制权威性", "Control authority", "输入或 horizon 的变化能否实质改变可行动作。"),
    ("Identifiability", "可辨识性", "Identifiability", "学生观测是否足以稳定区分教师的差异动作。"),
    ("Task-level Policy", "任务级策略", "Task-level policy", "为每个任务输出 defer 或数据中心分配动作。"),
    ("SLA", "服务等级协议", "Service-Level Agreement", "衡量等待或截止期违约的指标集合。"),
    ("Backlog", "积压队列", "Backlog", "尚未完成并等待后续调度的任务集合。"),
    ("Migration", "迁移", "Migration", "任务目标位置相对本地或原位置发生跨中心变化。"),
    ("Transmission", "传输成本", "Transmission cost", "跨数据中心任务迁移带来的数据传输代价。"),
    ("MAE", "平均绝对误差", "Mean Absolute Error", "预测值与真值绝对误差的平均。"),
    ("RMSE", "均方根误差", "Root Mean Squared Error", "对大误差更敏感的预测指标。"),
    ("Persistence", "持久性基线", "Persistence baseline", "直接用最近观测作为未来预测。"),
    ("P90/P95/P99", "分位点", "Percentiles", "分别表示 90%、95%、99% 样本不超过的阈值。"),
    ("Top-1", "首选动作准确率", "Top-1 accuracy", "概率最高动作与教师标签一致的比例。"),
    ("Top-2", "前二动作命中率", "Top-2 accuracy", "教师动作位于模型概率最高两个动作中的比例。"),
    ("Macro-F1", "宏平均 F1", "Macro-averaged F1", "各类别 F1 的等权平均。"),
    ("TRR", "教师相对恢复", "Teacher-relative recovery", "在 H1/Teacher 分歧样本上，BC 选择 Teacher 动作的比例。"),
    ("Fallback", "回退 H1", "Fallback to H1", "在分歧样本上 BC 仍选择 H1 动作的比例。"),
    ("Leakage", "信息泄漏", "Information leakage", "训练或控制使用了相应时刻本不应可见的信息。"),
    ("Timeline Contract", "时间轴契约", "Timeline contract", "预测索引、控制节点与未来分钟数的固定映射。"),
    ("Gate", "阶段门槛", "Decision gate", "达到预设证据标准后才进入下一阶段。"),
    ("Architecture A", "主架构 A", "Architecture A", "MPC 教师/基线→专家数据→BC→在线策略→评估。"),
    ("Architecture B", "候选架构 B", "Architecture B", "事件触发 MPC；当前保留为选择性优化机制。"),
    ("Architecture C", "候选架构 C", "Architecture C", "高层 RL + 下层 MPC；当前未实现。"),
]


def font(run, size=None, bold=None, color=None, latin="Times New Roman", east="DengXian"):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), latin); rf.set(qn("w:hAnsi"), latin); rf.set(qn("w:eastAsia"), east)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd"); node.set(qn("w:fill"), fill); pr.append(node)


def cell_margins(cell):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for key, value in (("top", 100), ("start", 120), ("bottom", 100), ("end", 120)):
        node = OxmlElement(f"w:{key}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); mar.append(node)
    pr.append(mar)


def table_geometry(table, widths):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(int(width / 2.54 * 1440))); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[min(i, len(widths)-1)]); cell_margins(cell)


def field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, txt, end]); return run


def bookmark(paragraph, name, idx):
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(idx)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(idx))
    paragraph._p.insert(0, start); paragraph._p.append(end)


def internal_link(paragraph, text, anchor):
    link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), anchor); link.set(qn("w:history"), "1")
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); col = OxmlElement("w:color"); col.set(qn("w:val"), BLUE); ul = OxmlElement("w:u"); ul.set(qn("w:val"), "single"); rpr.extend([col, ul]); run.append(rpr)
    txt = OxmlElement("w:t"); txt.text = text; run.append(txt); link.append(run); paragraph._p.append(link)


def configure(doc, header_text, numbered=True):
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2.35); sec.right_margin = Cm(2.15)
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5); normal._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
    normal.paragraph_format.line_spacing = 1.35; normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.first_line_indent = Pt(21); normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color in (("Heading 1", 17, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 11.5, TEAL)):
        st = doc.styles[name]; st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color); st._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian"); st.paragraph_format.keep_with_next = True; st.paragraph_format.space_before = Pt(10); st.paragraph_format.space_after = Pt(5)
    for name, size, color in (("Title", 28, NAVY), ("Subtitle", 12, GRAY)):
        st = doc.styles[name]; st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color); st._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
    doc.styles["Title"].font.bold = True
    for name, size, color in (("Lead", 11.5, NAVY), ("Figure Caption", 9, GRAY), ("Table Caption", 9, GRAY), ("Source Note", 8.5, GRAY), ("Formula", 11, NAVY), ("Callout", 10.5, NAVY)):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in doc.styles else doc.styles[name]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color); st._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian"); st.paragraph_format.first_line_indent = Pt(0)
    doc.styles["Figure Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; doc.styles["Formula"].font.name = "Cambria Math"; doc.styles["Formula"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp = sec.header.paragraphs[0]; r = hp.add_run(header_text); font(r, 8.5, True, NAVY)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = fp.add_run("ai_system_collaboration  |  "); font(r, 8, False, GRAY); font(field(fp, "PAGE", "1"), 8, False, GRAY)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); doc.settings.element.append(update)


def cover(doc, title, subtitle, lines, label):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(50); pr = p._p.get_or_add_pPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), NAVY); pr.append(sh); font(p.add_run(f"  {label}  "), 10, True, WHITE)
    p = doc.add_paragraph(style="Title"); font(p.add_run(title), 28, True, NAVY)
    p = doc.add_paragraph(style="Subtitle"); font(p.add_run(subtitle), 13, False, GRAY); p.paragraph_format.space_after = Pt(28)
    for line in lines:
        p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0); font(p.add_run(line), 10.5, False, NAVY)
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.space_before = Pt(70); font(p.add_run("证据截止：2026-09-01  |  本轮未新增训练或环境实验"), 9, False, GRAY)
    doc.add_page_break()


def toc(doc, title="目录", levels="1-3"):
    doc.add_heading(title, 1); p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0); field(p, 'TOC \\o "1-3" \\h \\z \\u'.replace("1-3", levels), "在 Word 中更新目录")
    p = doc.add_paragraph(style="Source Note"); p.add_run("首次打开请更新域，以刷新目录与页码。"); doc.add_page_break()


def callout(doc, label, text, fill=LIGHT, accent=TEAL):
    t = doc.add_table(rows=1, cols=1); table_geometry(t, [16.1]); c = t.cell(0, 0); shade(c, fill); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]; p.style = doc.styles["Callout"]; font(p.add_run(label + "  "), 10.5, True, accent); font(p.add_run(text), 10.5, False, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent = Cm(.65); p.paragraph_format.first_line_indent = Cm(-.35); p.paragraph_format.space_after = Pt(3); font(p.add_run(item), 10.5)


def table(doc, caption, headers, rows, widths, source=""):
    doc.add_paragraph(caption, style="Table Caption")
    t = doc.add_table(rows=1, cols=len(headers)); table_geometry(t, widths)
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, NAVY); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Pt(0); font(p.add_run(text), 9, True, WHITE)
    for ri, vals in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(vals):
            c = cells[i]; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2: shade(c, "F7F9FB")
            p = c.paragraphs[0]; p.paragraph_format.first_line_indent = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(str(val)) < 15 else WD_ALIGN_PARAGRAPH.LEFT; font(p.add_run(str(val)), 8.8)
    if source: doc.add_paragraph("来源：" + source, style="Source Note")


def figure(doc, idx, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Pt(0); shape = p.add_run().add_picture(str(FIGS[idx]), width=Cm(15.4)); shape._inline.docPr.set("title", f"图 {idx}"); shape._inline.docPr.set("descr", caption)
    doc.add_paragraph(f"图 {idx}  {caption}", style="Figure Caption"); doc.add_paragraph(f"数据：figure_data/fig{idx:02d}_*.csv；脚本：matlab/fig{idx:02d}_*.m", style="Source Note").alignment = WD_ALIGN_PARAGRAPH.CENTER


def para(doc, text, style=None):
    p = doc.add_paragraph(style=style); font(p.add_run(text), 10.5); return p


def add_core_terms(doc):
    doc.add_heading("核心术语定义与导航", 1)
    core = [x for x in TERMS if x[0] in {"MPC", "H1", "H4", "Oracle", "BC", "RL", "Transformer", "Triggered MPC", "Deployable", "Privileged Knowledge", "Identifiability", "Architecture A", "Architecture B", "Architecture C"}]
    anchors = {}
    for i, (abbr, cn, en, definition) in enumerate(core):
        p = doc.add_heading(f"{abbr}｜{cn}", 3); anchor = f"report_term_{i+1}"; bookmark(p, anchor, i+1); anchors[abbr] = anchor
        para(doc, f"{en}。{definition}")
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0); font(p.add_run("快速跳转："), 9, True, GRAY)
    for i, (term, anchor) in enumerate(anchors.items()):
        if i: p.add_run(" · ")
        internal_link(p, term, anchor)
    doc.add_page_break(); return anchors


def h1(doc, number, title): doc.add_heading(f"{number} {title}", 1)
def h2(doc, number, title): doc.add_heading(f"{number} {title}", 2)


def build_report():
    d = Document(); configure(d, "完整研究工作汇报")
    cover(d, "多数据中心 AI 调度：完整研究工作汇报", "从真实性审计到 Forecast + Privileged MPC Teacher + Deployable Neural Policy", ["项目目录：D:\\ai_system_collaboration", "Git：feature/baseline-repair-information-contract-v1 @ 90eb972f78742603b522fabc3e1cf65391434418", "SustainCluster：3f6ea95cb835b89ba50b0ef76d66d14b8037643e（clean）", "材料性质：阶段研究总结、技术路线复盘与下一阶段 Gate 设计"], "FINAL RESEARCH REPORT v1")
    toc(d); anchors = add_core_terms(d)

    h1(d, 1, "研究背景、目标与边界"); h2(d, "1.1", "研究问题")
    para(d, "本项目研究多数据中心任务调度：在异构算力、动态工作负载、电价、碳强度、容量、截止期和跨中心传输共同存在时，控制器需要决定任务应延迟还是分配到某个数据中心。目标不是追求单一离线分数，而是建立信息可部署、时间语义正确、结果可复核的完整链路。")
    para(d, "阶段路线依次经历 workload 与接口审计、Forecast Dataset、Transformer、forecast-aware MPC、控制权威性、formulation repair、Triggered MPC、Expert Dataset v2 和 BC v2。负结果被保留，因为它们实际推动了研究问题收敛。")
    callout(d, "当前主判断", "路线已从笼统的 Online RL + MPC 收敛为 Forecast + Privileged MPC Teacher + Deployable Neural Policy；RL 是后续 Gate，不是当前已验证成果。", PALE_GOLD, GOLD)
    figure(d, 1, "研究问题到当前路线：审计、修复、验证和角色决策")
    para(d, "如图 1 所示，本报告按“问题→假设→实验→发现→修正→新问题”组织；本轮不新增训练、不重跑环境实验，也不把旧口径与修复后口径混合。")
    table(d, "表 1  当前版本与工作边界", ["项目", "状态", "解释"], [["Git HEAD", "90eb972f…", "证据冻结点；工作区包含此前阶段未提交改动"], ["核心代码", "未修改", "只新增报告脚本与 final_research_report_v1"], ["训练/实验", "未执行", "沿用已有 artifact 和测试"], ["第三方源码", "未修改", "SustainCluster 独立仓库 clean"]], [3.1, 2.5, 10.5], "Git 与任务约束")

    h1(d, 2, "系统定义、状态与动作空间"); h2(d, "2.1", "SustainCluster 与 Adapter")
    para(d, "SustainCluster 提供多数据中心任务、资源、执行和评价环境；项目侧 Adapter 把第三方字段映射为稳定状态和动作语义。环境可持有真实运行时与未来轨迹，但控制器输入必须由 Information Contract 授权。")
    figure(d, 2, "多数据中心热电耦合调度框架")
    para(d, "如图 2 所示，任务流经 Adapter 后，由 MPC、BC 或后续 RL 输出任务级动作，再由环境统计 SLA、能耗、碳、迁移、传输和 backlog。")
    h2(d, "2.2", "34 维观测与 6 类语义动作"); para(d, "学生策略使用 34 维可部署观测。动作 0 为 defer，动作 1–5 对应五个语义数据中心；可行集合随容量、等待与 deadline 变化。")
    figure(d, 3, "任务级语义动作空间")
    para(d, "如图 3 所示，defer 表示进入等待队列而非丢弃；修复后等待计数每个 step 恰好增长一次。")

    h1(d, 3, "主线与 Architecture A/B/C 定位")
    h2(d, "3.1", "Architecture A：当前主路径"); para(d, "SustainCluster→Adapter→MPC Expert/Baseline→Expert Dataset→BC→Online Policy→Evaluation。MPC 同时是教师和优化基线；BC 是可部署神经策略候选；RL 只保留 BC→RL 接口。")
    h2(d, "3.2", "Architecture B/C：候选路径"); para(d, "Architecture B 使用 deployable 风险触发 H4；现有证据表明分歧富集有效但在线收益有限。Architecture C 为高层 RL + 下层 MPC，目前没有实现或实验，不进入第一版主路径。")
    table(d, "表 2  三类架构定位", ["架构", "角色", "证据", "下一步"], [["A", "主架构", "Teacher、Expert Dataset、BC 均有阶段证据", "先解决可辨识性"], ["B", "候选机制", "P95 触发 5.625%，收益有限", "选择性优化/采样"], ["C", "远期候选", "只有架构论证", "待 Gate 后评估"]], [1.5, 3.0, 7.0, 4.6], "architecture audit；triggered_mpc_v1")

    h1(d, 4, "历史 MPC、BC 与 SAC")
    h2(d, "4.1", "旧 BC 与闭环"); para(d, "审计前旧 BC 得到 overall accuracy 96.981%、assign accuracy 90.492%、migration F1 94.790%，并完成真实多动作闭环。其 completed 与 SLA 接近 MPC、推理时延更低、illegal/overflow 为 0；但 reward、迁移和 transmission 并未全面优于 MPC。")
    para(d, "旧结果证明专家动作可被神经网络压缩，但旧专家数据、状态信息与 formulation 不同，不能用于覆盖 BC v2 的新结果。`original_rule_local_only` 的 reward/SLA 因接口限制不可解释。")
    h2(d, "4.2", "BC-init SAC 的起点改善与洗脱"); para(d, "历史公平比较中 BC-init 初始 reward -1821.22，random-init -3641.44；10,000 步后 BC-init reward 降至 -2432.21，SLA 增至 2080，expert agreement 从约 0.9256 降至 0.7690。Critic warmup 与 behavior regularization 小实验仍未通过 Gate。")
    callout(d, "边界", "可否定的是当前 reward + task-level Critic + vanilla discrete SAC 配方，不能泛化为所有 RL 无效。", LIGHT, RED)

    h1(d, 5, "Workload 真实性与泄漏审计")
    h2(d, "5.1", "49 天重复"); para(d, "本地 workload pickle 有 37,552 行，外层时间 2020-01-01 至 2021-01-26；15 分钟网格预期 37,632 行，缺失 80。相邻 49 天块共 32,858 个对齐区间，任务矩阵逐元素 exact match 100%，任务数相关系数 1。")
    figure(d, 4, "Workload 49 天重复审计")
    para(d, "如图 4 所示，外层近全年不等于全年独立样本。Forecast Dataset v1 只使用首个已验证唯一周期，避免随机切分重复块造成泄漏。")
    h2(d, "5.2", "时间轴、duration 与 Alibaba 2026"); para(d, "任务内部 processing timeline 为 1970-01-26 至 1970-03-15。true_duration 来自完成后的实现值，只能作模拟器真值；控制器使用 estimated_duration。Alibaba 2026 方向已评估，但本地未形成可复核的 15 分钟统一事实表，且接入与许可复杂，本阶段暂缓。")

    h1(d, 6, "Baseline Repair 与信息契约")
    para(d, "旧适配器把索引 8 的 avg_gpu_wrk_mem 当作网络带宽；修复后 bandwidth 使用索引 9 的 bandwidth_gb。该错误会影响迁移与传输计算，相关历史数值只能按旧口径保留。")
    table(d, "表 3  控制器信息契约", ["类别", "字段", "在线可见", "用途"], [["当前状态", "资源占用、当前价格/碳、等待", "是", "H1/BC/Trigger"], ["外部未来", "未来价格、未来碳", "可由服务提供", "deployable H4"], ["预测负载", "未来 CPU/GPU/内存", "模型输出", "forecast-aware MPC"], ["环境真值", "true_duration、真实未来 workload", "否", "Oracle Teacher/上界"]], [3.0, 5.0, 3.0, 5.1], "baseline_repair_information_contract_v1")
    para(d, "Oracle future workload 只允许离线 Teacher 使用，Student 输入保持 deployable；两者差异构成 privileged-knowledge distillation 的研究对象。")

    h1(d, 7, "Forecast Dataset v1")
    para(d, "数据集使用首个唯一 49 天周期，共 4,704 个 15 分钟网格；4,694 步非空，10 个缺失网格补零。目标为未来 task count、CPU、GPU、memory demand；输入为四类负载历史与四类日历特征。")
    figure(d, 5, "Forecast Dataset v1 的时间切分与滑动窗口")
    para(d, "如图 5 所示，history=96 步（24 小时），horizon=4（+15/+30/+45/+60 分钟）；Train/Validation/Test 为 35/7/7 天，对应 3,261/573/573 个窗口，scaler 只拟合 Train。")
    bullets(d, ["不使用重复块扩充样本。", "不随机打乱时间切分。", "窗口历史端点早于预测端点。", "manifest 记录数据哈希、列、split 与 scaler。"])

    h1(d, 8, "Transformer Forecast v1")
    para(d, "Transformer 输入 [N,96,8]、输出 [N,4,4]；d_model=64、4 heads、2 encoder layers、feedforward=256、dropout=0.1、head=128，总参数 110,928。")
    figure(d, 6, "Transformer Forecast v1 模型结构")
    para(d, "推荐 seed 33，最佳 epoch 32，validation loss 0.96577248；checkpoint SHA-256 为 9CD85A0D229E135106EA1E49CD8FEC8936D7B6349DA41531564E565F063569E2。batch=1 CPU 推理均值 1.812 ms、P95 2.504 ms。")
    figure(d, 7, "GPU 需求预测的平均 MAE")
    para(d, "如图 7 所示，四个 horizon 的 GPU MAE 均低于 Persistence，+60 分钟从 74.79 降至 52.65，改善 29.61%。")
    figure(d, 8, "GPU 需求预测 RMSE 与 P90 峰值短板")
    para(d, "如图 8 所示，总体 RMSE 改善，但 P90 峰值 Transformer RMSE 325.54，高于 Persistence 308.54；结论为 MIXED，峰值与不确定性是下一阶段重点。")

    h1(d, 9, "Forecast-aware MPC v1")
    para(d, "Forecast-aware MPC 把 Transformer/Persistence 的未来负载映射为数据中心未来压力，并与未来价格/碳共同进入 H4。比较严格区分 H1、H4 Oracle、H4 Persistence 与 H4 Transformer。")
    para(d, "早期现象是预测误差改善但 H4 Transformer 与 H1 动作高度一致。该结果没有直接否定 Forecast，而是提出新的因果问题：future signal 是否正确进入 state/objective/action，真实 trace 是否足以改变动作。")
    callout(d, "关键认识", "Forecast accuracy 不自动等于 control benefit；必须单独证明 forecast→state→objective→action 链。")

    h1(d, 10, "MPC Control Authority 与 Future Sensitivity")
    para(d, "合成容量释放、低价和 GPU burst 场景证明控制器在足够压力下会改变 defer/execute 与目标数据中心，因此接口并非完全失效。")
    figure(d, 10, "未来压力与动作切换：合成响应与真实低压力")
    para(d, "如图 10 所示，原始敏感性 sweep 在 ALL_RESOURCES forecast_scale=3.0 时将目标从 DC1 切换到 DC2；真实 Transformer GPU forecast/capacity 的 P50/P90/P99 仅 1.175%/3.957%/7.037%，远离切换区。根因为 accounting、objective、timeline、低压力 trace 和时间语义的组合。")

    h1(d, 11, "MPC Formulation Repair v1")
    para(d, "旧 H4 节点为 [t,+15,+30,+45]，名义 horizon=4 但真实只到 +45 分钟；修复后内部节点为 [t,+15,+30,+45,+60]，四个预测值映射到未来节点 1–4。waiting 只在单一位置增长，defer 序列为 0→1→2。")
    figure(d, 9, "Old H4 与 repaired H4 的时间轴契约")
    para(d, "如图 9 所示，修复恢复了 forecast[k] 与 t+15k 的固定语义。")
    figure(d, 11, "Repaired H1 与 H4 的核心指标")
    para(d, "如图 11 所示，H4 Oracle reward -1639.04 高于 H1 -1655.40、electricity 略低，但 SLA 1234.6 高于 H1 1224.6。H4 Transformer reward -1654.86，与 H1 接近，求解约 4.61 ms，H1 约 1.99 ms。H4 只能描述为小幅、条件性收益。")

    h1(d, 12, "Triggered MPC v1")
    para(d, "Triggered MPC 以当前与 forecast occupancy 的最大值构造 deployable risk；P95 阈值 0.5854637688。未触发走 H1，触发才调用 H4。")
    figure(d, 12, "Triggered MPC 的 deployable 风险门控")
    para(d, "P95 下 960 步中 H4 调用 54 次、H1 调用 906 次，H4 调用量减少 94.375%。")
    figure(d, 13, "触发状态中的 H1/H4 分歧富集")
    para(d, "如图 13 所示，P95 触发状态分歧 10.51%，非触发 0.82%，但相对 H1 的 reward 增益仅 +0.326，stage cost 改善约 1.201。其价值更适合解释为选择性优化和难例采样，而非强在线收益。")

    h1(d, 13, "Repaired MPC Expert Dataset v2")
    para(d, "Teacher 使用 repaired H4 Oracle，Student 只保留当前 deployable 34 维观测。80 episodes 共 7,680 steps、259,920 task rows、259,397 unique task IDs；按 episode/seed 划分 Train 181,944、Validation 38,988、Test 38,988。")
    figure(d, 14, "Repaired MPC Expert Dataset v2 结构")
    para(d, "数据 SHA-256 为 27E0B22558E1F00AA22E3B899D505E7506C10192ED99EA360278C8D819FF9E3C，solver failures=0。Oracle 只产生标签，不进入学生输入。")
    figure(d, 15, "H1 与 Oracle Teacher 的任务级和高风险分歧")
    para(d, "如图 15 所示，总体任务级分歧 3.68%，state-level any disagreement 16.82%；高风险任务分歧 19.35%，低风险 3.07%。真正需要学习的是稀疏、未来敏感的关键差异。")

    h1(d, 14, "BC v2：Offline Privileged-Knowledge Distillation")
    para(d, "ActorNet 为 34→256→256→6，LayerNorm+ReLU，共 77,318 参数。seeds 11/22/33 最佳 epoch 为 6/9/4，推荐 seed 22 的 validation loss 0.51651668。")
    figure(d, 16, "BC v2 与 H1 Teacher 的离线 Top-1 准确率")
    para(d, "如图 16 所示，H1 对 Teacher 的准确率 96.79%；BC 三种子 76.50%/73.17%/74.93%，均值 74.86%±1.36%；Top-2 均值 90.99%，macro-F1 0.5499。一般模仿弱于 H1。")
    figure(d, 17, "BC v2 在 H1/Teacher 分歧子集上的恢复")
    para(d, "如图 17 所示，1,252 个分歧样本上的 TRR 均值 34.66%，H1 fallback 50.53%，其他错误 14.80%；高风险分歧 TRR 36.55%。195 个分歧状态中 any recovery 约 57.78%、full recovery 29.06%；BC defer predictions=0。结论为 GENERAL IMITATION WEAK，不具备闭环上线条件。")

    h1(d, 15, "当前洞见与证据等级")
    table(d, "表 4  当前证据边界", ["等级", "内容", "表达规则"], [["已经证实", "49天重复；字段/Oracle/timeline bug；平均预测改善；触发富集；BC v2 模仿弱", "可作事实并附来源"], ["目前倾向", "在线 H4 价值有限；MPC 更适合作为 Teacher；BC 问题与可辨识性有关", "使用倾向/当前判断"], ["尚未证实", "不确定性预测改善控制；分歧蒸馏过 Gate；BC→RL 保持安全", "进入最小判别实验"], ["不能声称", "BC v2 已部署；RL 已优于基线；H4 全面优于 H1；392天为独立全年；C 已验证", "从标题与结论排除"]], [2.4, 9.3, 4.4], "全部当前 artifacts")
    para(d, "Teacher/H1 平均分歧稀疏，但高风险分歧集中；Student 又看不到 Teacher 的未来真值。普通交叉熵被大量共识样本主导，因此下一问题是 deployable observation 的可辨识性，而非直接增加训练步数。")
    callout(d, "Architecture A Gate", "保留 Adapter、Teacher、Expert Dataset 与 deployable policy；暂停直接 BC→SAC。离线关键差异、闭环安全和多种子稳定性同时成立后再恢复 RL。", PALE_GOLD, GOLD)

    h1(d, 16, "下一阶段路线、预期结果与结论")
    h2(d, "16.1", "近期优先级"); bullets(d, ["P0：disagreement-aware identifiability audit，估计可辨识上界与 risk-conditioned entropy。", "P0：比较 weighted/focal/pairwise/selective distillation，首看 high-risk TRR 与 calibration。", "P0：建立闭环 Gate：illegal/overflow=0、SLA 不劣于 H1、关键差异多种子稳定。", "P1：增加峰值加权、分位数或 ensemble，不确定性必须通过控制收益验证。", "P1：把 Triggered MPC 用于难例采样和主动标注。", "P2：Gate 通过后再比较 frozen BC、regularized BC→RL 与 selective correction。"])
    table(d, "表 5  最小判别实验", ["实验", "变量", "指标", "决策"], [["Identifiability", "deployable features vs privileged label", "条件熵/上界", "不足则改观测"], ["Weighted distillation", "共识/分歧/风险权重", "TRR/fallback/F1", "优于 H1 再闭环"], ["Selective policy", "BC confidence+H1 fallback", "coverage-risk/SLA", "优先安全"], ["Uncertainty forecast", "quantile/ensemble/peak loss", "peak error/trigger capture", "证明预测→控制"], ["BC→RL Gate", "frozen/regularized/selective", "reward/SLA/agreement", "不只看初始 reward"]], [3.0, 5.0, 4.6, 3.5], "当前失败模式推导；尚未执行")
    h2(d, "16.2", "限制"); bullets(d, ["只覆盖一个已验证唯一 49 天周期，不能代表全年季节变化。", "Transformer 峰值性能弱，真实 GPU 压力远低于容量。", "Teacher 使用 Oracle future workload，Student 存在结构性 privileged gap。", "BC v2 尚无闭环证据。", "SustainCluster、Alibaba workload 与 checkpoint 公开发布许可仍需确认。", "当前工作区有此前阶段未提交改动，报告不是新冻结 commit。"])
    figure(d, 18, "当前决策、下一阶段 Gate 与 RL 恢复条件")
    para(d, "如图 18 所示，下一阶段以可辨识性与 disagreement-aware distillation 为主，Forecast 不确定性并行推进。RL 的恢复条件是离线关键差异、闭环安全和多种子稳定性三类证据同时成立。")
    h2(d, "16.3", "结论"); para(d, "本阶段建立了可信的数据、信息与时间语义：识别重复与 Oracle 泄漏，修复 bandwidth、waiting 和 H4 timeline，证明 Transformer 的平均预测价值，厘清 H4 在线收益边界，构建 privileged Teacher 数据集，并由 BC v2 暴露可辨识性问题。当前最稳健主线是 Forecast + Privileged MPC Teacher + Deployable Neural Policy。")
    d.add_page_break(); d.add_heading("证据索引", 1)
    table(d, "表 6  主要 artifact", ["主题", "路径"], [["Workload audit", "artifacts/workload_information_audit_v1/"], ["Information contract", "artifacts/baseline_repair_information_contract_v1/"], ["Forecast Dataset", "artifacts/forecast_dataset_v1/"], ["Transformer", "artifacts/transformer_forecast_v1/"], ["MPC authority", "artifacts/mpc_control_authority_diagnosis_v1/"], ["MPC repair", "artifacts/mpc_formulation_repair_v1/"], ["Triggered MPC", "artifacts/triggered_mpc_v1/"], ["Expert Dataset v2", "artifacts/repaired_mpc_expert_dataset_v2/"], ["BC v2", "artifacts/bc_v2_offline/"], ["图件", "artifacts/final_research_report_v1/figure_manifest.csv"]], [4.2, 11.9], "source_manifest.md")
    d.save(REPORT)


SLIDES = [
    (1, "封面", "这项研究最终解决什么问题？", ["多数据中心热电耦合调度", "两日研究链与路线重构", "证据截止 2026-09-01"], None, "不是单点算法汇报，而是证据驱动的路线收敛。"),
    (2, "研究路线总览", "为什么路线与最初设想不同？", ["审计→预测→控制→教师→学生", "负结果推动角色变化"], 1, "Forecast、MPC、BC/RL 的角色由实验决定。"),
    (3, "系统与接口", "算法实际控制什么？", ["SustainCluster", "Adapter", "任务级策略与评价"], 2, "先固定系统语义，再比较算法。"),
    (4, "动作空间", "策略每步输出什么？", ["0=defer", "1–5=五个数据中心", "可行性由约束决定"], 3, "MPC 与 BC 共享动作语义。"),
    (5, "历史 Architecture A", "为什么 MPC→BC→SAC 曾可行？", ["旧 BC 96.981%", "完成多动作闭环", "BC 初始化改善 SAC 起点"], None, "历史正证据不保证在线保持。"),
    (6, "BC-init SAC 洗脱", "更好起点为何没有更好终点？", ["初始 -1821.22 vs -3641.44", "10k 后 reward/SLA/agreement 恶化", "Warmup/KL 未过 Gate"], None, "否定当前配方，不是否定所有 RL。"),
    (7, "Workload 审计", "是否真的有全年独立负载？", ["37,552 行", "49 天逐元素重复", "Forecast 只取唯一周期"], 4, "主动缩小数据范围以避免泄漏。"),
    (8, "信息契约", "哪些字段在线可见？", ["bandwidth 修复", "true vs estimated duration", "Oracle vs deployable"], None, "先修信息边界再解释收益。"),
    (9, "Forecast Dataset", "如何避免时间泄漏？", ["49 天/4,704 步", "96→4", "35/7/7 天", "train-only scaler"], 5, "时间顺序是数据集核心。"),
    (10, "Transformer", "预测器规模与推理是否可用？", ["[N,96,8]→[N,4,4]", "110,928 参数", "1.812 ms"], 6, "模型轻量，关键看相对基线。"),
    (11, "平均预测结果", "正常场景改善多少？", ["四个 horizon MAE 全改善", "+60 min 改善 29.61%"], 7, "平均预测价值已证实。"),
    (12, "峰值短板", "平均更好是否等于峰值更好？", ["P90 RMSE 325.54 vs 308.54", "结论 MIXED"], 8, "峰值决定后续不确定性研究。"),
    (13, "H4 时间轴修复", "名义 H4 是否到 +60？", ["旧版只到 +45", "修复到 +60", "waiting 单次增长"], 9, "先修 formulation 再比较。"),
    (14, "Control Authority", "无收益是接口失效还是压力不足？", ["ALL_RESOURCES scale=3 时切换 DC", "真实 P90 仅 3.957%"], 10, "算法可响应，真实 trace 太轻。"),
    (15, "Repaired H1/H4", "H4 是否全面胜出？", ["reward/能耗略优", "SLA 略差", "求解约 2.3 倍"], 11, "H4 只有条件性小收益。"),
    (16, "Triggered MPC", "能否只在困难状态调用 H4？", ["P95=0.5855", "H4 54/960 次"], 12, "Trigger 集中优化计算。"),
    (17, "触发分歧价值", "触发状态是否更难？", ["触发分歧 10.51%", "非触发 0.82%", "reward +0.326"], 13, "分歧富集成立，在线收益有限。"),
    (18, "Expert Dataset v2", "如何合法使用未来真值？", ["H4 Oracle Teacher", "34 维 deployable Student", "259,920 rows"], 14, "Oracle 只产生标签。"),
    (19, "Teacher/H1 分歧", "需要学习的差异多集中？", ["总体 3.68%", "高风险 19.35%", "state any 16.82%"], 15, "关键差异稀疏而集中。"),
    (20, "BC v2 总体结果", "普通 BC 能提取 privileged knowledge 吗？", ["H1 96.79%", "BC 74.86%±1.36%", "Top-2 90.99%"], 16, "一般模仿弱。"),
    (21, "关键分歧恢复", "BC 是否学会 Teacher 相对 H1 的改进？", ["TRR 34.66%", "fallback 50.53%", "defer=0"], 17, "privileged knowledge 未稳定迁移。"),
    (22, "决策与下一步", "现在做什么、暂停什么？", ["Identifiability", "分歧感知蒸馏", "峰值/不确定性", "Gate 后再 RL"], 18, "当前路线：Forecast + Teacher + Deployable Policy。"),
]


def build_outline():
    d = Document(); configure(d, "汇报PPT详细提纲", numbered=False)
    cover(d, "汇报 PPT 详细提纲", "多数据中心 AI 调度：从真实性审计到 privileged Teacher", ["建议：22 页，30–40 分钟", "听众：项目组、导师、协作成员", "所有数值以本地 artifact 为准", "所有图件均由 MATLAB 从 CSV 生成"], "PRESENTATION OUTLINE")
    toc(d, "提纲导航", "1-1")
    for idx, title, question, body, fig_idx, key in SLIDES:
        d.add_heading(f"第 {idx:02d} 页｜{title}", 1); callout(d, "本页回答", question, LIGHT, BLUE)
        d.add_heading("页面正文", 2); bullets(d, body)
        d.add_heading("建议图件与路径", 2)
        if fig_idx:
            para(d, f"图 {fig_idx}：{FIGS[fig_idx].name}\nartifacts/final_research_report_v1/figures/{FIGS[fig_idx].name}")
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Pt(0); shape = p.add_run().add_picture(str(FIGS[fig_idx]), width=Cm(9.3)); shape._inline.docPr.set("descr", f"第{idx}页建议图")
        else: para(d, "无强制图件；使用极简事实对比，不新增装饰性图片。")
        d.add_heading("关键讲点", 2); para(d, key)
        d.add_heading("口头讲稿", 2); para(d, f"这一页要让听众先回答“{question}”。围绕页面事实解释证据的适用边界，强调：{key} 不把尚未完成的实验写成结论。")
        d.add_heading("转场", 2); para(d, "沿着“问题→证据→下一问题”进入下一页，避免只罗列数字。" if idx < 22 else "结束汇报，进入问题讨论。")
        if idx < 22: d.add_page_break()
    d.add_page_break(); d.add_heading("统一视觉与数据规范", 1)
    table(d, "表 1  PPT 规范", ["项目", "规范"], [["画幅", "16:9；正文安全边距 5%"], ["字体", "中文等线；英文/数字 Times New Roman；公式 Cambria Math"], ["颜色", "深蓝主色；青绿=有效；金色=决策；红色=风险"], ["标题", "每页回答一个问题，最多两行"], ["来源", "页脚写 artifact/CSV；历史结果标注旧口径"], ["结论", "明确已证实/倾向/未证实/不能声称"]], [3.5, 12.6], "figure_manifest.csv 与用户规范")
    d.save(OUTLINE)


def build_glossary():
    d = Document(); configure(d, "术语与符号说明", numbered=False)
    cover(d, "术语与符号说明", "MPC、Forecast、BC/RL 与证据边界的统一口径", ["适用：完整研究报告、PPT 提纲与正式汇报", "用途：统一中英文、公式符号与禁止混用项", "目录与术语索引支持内部跳转"], "GLOSSARY & NOTATION")
    toc(d, levels="1-1"); d.add_heading("术语索引", 1); anchors = {abbr: f"glossary_{i+1}" for i, (abbr, _, _, _) in enumerate(TERMS)}
    p = d.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    for i, (term, anchor) in enumerate(anchors.items()):
        if i: p.add_run(" · ")
        internal_link(p, term, anchor)
    d.add_page_break(); d.add_heading("核心术语", 1)
    warnings = {"Oracle": "不得与 deployable forecast 混写。", "BC": "BC v2 只有离线证据；旧 BC 属于不同口径。", "RL": "本阶段未训练，不能声称优于基线。", "Architecture C": "当前没有实现或实验。"}
    for i, (abbr, cn, en, definition) in enumerate(TERMS):
        p = d.add_heading(f"{abbr}｜{cn}", 2); bookmark(p, anchors[abbr], 500+i)
        para(d, f"英文：{en}", "Lead"); para(d, definition)
        if abbr in warnings: callout(d, "口径提醒", warnings[abbr], PALE_GOLD, GOLD)
    d.add_heading("符号与公式", 1)
    table(d, "表 1  常用符号", ["符号", "含义", "范围"], [["t", "当前 15 分钟控制步", "step"], ["H", "预测/控制 horizon", "1 或 4"], ["x_t", "可部署状态", "Student 34 维"], ["a_t", "任务级动作", "{0,…,5}"], ["ŵ_{t+k}", "预测 workload", "k=1,…,4"], ["w*_{t+k}", "Oracle workload", "Teacher only"], ["J_H", "H 步 MPC 目标", "stage cost"], ["r_t", "环境 reward", "越高越好"], ["τ", "Trigger 阈值", "P95=0.5854637688"], ["π_BC", "BC 策略", "6 类概率"], ["y*", "Teacher 标签", "{0,…,5}"]], [2.5, 9.0, 4.6], "配置、manifest 与报告")
    para(d, "J_H = Σ_{k=0}^{H} ℓ(x_{t+k}, a_{t+k}; price, carbon, SLA, migration, transmission)", "Formula")
    para(d, "risk_t = max_d,k (current_load_{d,t} + forecast_load_{d,t+k}) / capacity_d", "Formula")
    para(d, "TRR = P(π_BC(x)=y* | y_H1 ≠ y*)", "Formula")
    d.add_heading("必须分开的口径", 1)
    table(d, "表 2  易混概念", ["概念 A", "概念 B", "区别"], [["Oracle H4", "Transformer/Persistence H4", "未来真值 Teacher vs 可部署"], ["Old H4", "Repaired H4", "+45 min vs +60 min"], ["旧 BC", "BC v2", "数据、标签和问题不同"], ["Forecast accuracy", "Control benefit", "误差改善不保证控制收益"], ["Task disagreement", "State any disagreement", "任务比例 vs 状态内任一分歧"], ["SAC reward", "MPC objective", "项和权重并不完全一致"]], [4.0, 4.2, 7.9], "source_manifest.md")
    d.save(GLOSSARY_DOC)


def audit_docx(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8"); styles = z.read("word/styles.xml").decode("utf-8"); settings = z.read("word/settings.xml").decode("utf-8")
    return {"file": path.name, "bytes": path.stat().st_size, "figures": xml.count("<wp:inline"), "tables": xml.count("<w:tbl>"), "bookmarks": xml.count("bookmarkStart"), "internal_links": xml.count("w:anchor="), "TOC": "TOC \\o" in xml, "PAGE": "PAGE" in xml, "updateFields": "updateFields" in settings, "DengXian": "DengXian" in styles or "DengXian" in xml, "TimesNewRoman": "Times New Roman" in styles or "Times New Roman" in xml, "CambriaMath": "Cambria Math" in styles or "Cambria Math" in xml}


def main():
    if len(FIGS) != 18: raise RuntimeError(f"Expected 18 figures, found {len(FIGS)}")
    build_report(); build_outline(); build_glossary()
    audits = [audit_docx(p) for p in (REPORT, OUTLINE, GLOSSARY_DOC)]
    manifest = {"documents": audits, "figures": {"png": len(list(FIGDIR.glob('Fig*.png'))), "fig": len(list((OUT/'figures_matlab').glob('Fig*.fig'))), "matlab": len(list((OUT/'matlab').glob('fig*.m'))), "csv": len(list((OUT/'figure_data').glob('fig*.csv')))}, "scope": {"core_code_modified": False, "models_retrained": False, "dataset_regenerated": False, "commit": False, "push": False}, "git_head": "90eb972f78742603b522fabc3e1cf65391434418", "sustaincluster_head": "3f6ea95cb835b89ba50b0ef76d66d14b8037643e"}
    (OUT / "report_generation_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# Report Generation Manifest", "", *[f"- `{x['file']}`: figures={x['figures']}, tables={x['tables']}, bookmarks={x['bookmarks']}, internal_links={x['internal_links']}, TOC={x['TOC']}" for x in audits], "", "## Figures", "", f"- PNG: {manifest['figures']['png']}", f"- MATLAB .fig: {manifest['figures']['fig']}", f"- MATLAB scripts: {manifest['figures']['matlab']}", f"- Raw CSV: {manifest['figures']['csv']}", "", "## Scope", "", "- Core code modified: NO", "- Models retrained: NO", "- Dataset regenerated: NO", "- Commit: NO", "- Push: NO", "", "Render page counts and visual QA are appended by validation."]
    (OUT / "report_generation_manifest.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
